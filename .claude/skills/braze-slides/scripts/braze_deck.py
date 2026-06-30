# -*- coding: utf-8 -*-
"""
braze_deck.py — reusable BrazeAI-styled slide builder (python-pptx).

Brand values are extracted from the official Braze "Decisioning Studio"
slideware theme. Output is a 10 x 5.625in (Braze 16:9) .pptx using the
Braze palette, Helvetica Neue, and the real BrazeAI logo top-right.

Two ways to use it:
  1. Archetype builders (title/section/two_column/cards/stat_row/table/
     bar_chart/callout/closing) — the fast path for turning an outline or
     markdown into on-brand slides.
  2. A blank slide + the exposed primitives (box/para/card/rect/divider/
     pill/logo) — for bespoke slides that still inherit the brand.

Author a small build script that imports BrazeDeck and calls these; render
and eyeball it (in Cowork), then iterate. See example_blc_deck.py.
"""
import os
from pptx import Presentation
from pptx.util import Inches as _IN, Pt as _PT, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ----------------------------------------------------------------------------
# Canvas: design in a 13.333 x 7.5 space, emit at 0.75 -> Braze's 10 x 5.625in.
# Author everything in "design inches"; SCALE keeps proportions identical while
# matching Braze's real slide size so these slides drop straight into a deck.
# ----------------------------------------------------------------------------
SCALE = 0.75
def IN(v): return _IN(v * SCALE)
def PT(v): return _PT(v * SCALE)

DW, DH = 13.333, 7.5          # design width/height
ML = 0.6                       # left/right margin (design in)
CW = DW - 2 * ML               # content width

# ----------------------------------------------------------------------------
# Brand kit (official Braze theme)
# ----------------------------------------------------------------------------
PURPLE   = RGBColor(0x80, 0x1E, 0xD7)   # accent2 — primary brand purple
PURPLE_D = RGBColor(0x30, 0x02, 0x66)   # dk2 — deep purple
GOLD     = RGBColor(0xFF, 0xA5, 0x24)   # accent5 — gold/amber
BERRY    = RGBColor(0x91, 0x18, 0x6E)   # accent3
RED      = RGBColor(0xE9, 0x37, 0x1F)   # accent4 (use sparingly)
LAV      = RGBColor(0xEF, 0xED, 0xFF)   # lt2 — light lavender (card fills)
PERI     = RGBColor(0xC9, 0xC4, 0xFF)   # accent6 — periwinkle
INK      = RGBColor(0x1A, 0x1A, 0x24)   # body text
SLATE    = RGBColor(0x55, 0x52, 0x66)   # secondary text
MUTE     = RGBColor(0x8A, 0x86, 0x99)   # tertiary / captions
LINE     = RGBColor(0xE3, 0xDF, 0xF5)   # hairlines / borders
ONLAV    = RGBColor(0xD9, 0xD3, 0xFB)   # light text on purple fills
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
PAPER    = RGBColor(0xFA, 0xF9, 0xFE)

FONT = "Helvetica Neue"        # leadlines Bold; Arial is the safe fallback
LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "..", "assets", "brazeai_logo.png")
LOGO_AR = 1200 / 513           # native logo aspect ratio


class BrazeDeck:
    def __init__(self, logo_path=None):
        self.prs = Presentation()
        self.prs.slide_width = IN(DW)
        self.prs.slide_height = IN(DH)
        self._blank = self.prs.slide_layouts[6]
        self.logo_path = logo_path or LOGO

    # ----- low-level primitives (usable on any slide) -----
    def blank(self):
        """Add a blank slide (white). Returns the slide for bespoke work."""
        s = self.prs.slides.add_slide(self._blank)
        self.rect(s, 0, 0, DW, DH, WHITE)
        return s

    def box(self, s, l, t, w, h, anchor=None):
        b = s.shapes.add_textbox(IN(l), IN(t), IN(w), IN(h))
        tf = b.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor: tf.vertical_anchor = anchor
        return b, tf

    def para(self, tf, runs, align=PP_ALIGN.LEFT, first=False,
             before=None, after=None, line=None):
        """runs: a dict or list of dicts: {text,size,color,bold,italic,font,spc,super}."""
        if first and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if before is not None: p.space_before = PT(before)
        if after is not None:  p.space_after = PT(after)
        if line is not None:   p.line_spacing = line
        if isinstance(runs, dict): runs = [runs]
        for sp in runs:
            r = p.add_run(); r.text = sp['text']
            r.font.size = PT(sp.get('size', 12))
            r.font.bold = sp.get('bold', False)
            r.font.italic = sp.get('italic', False)
            r.font.name = sp.get('font', FONT)
            r.font.color.rgb = sp.get('color', INK)
            if 'spc' in sp: r.font._rPr.set('spc', str(sp['spc']))
            if sp.get('super'): r.font._rPr.set('baseline', '30000')
        return p

    def _shape(self, s, kind, l, t, w, h, fill=None, line_c=None, line_w=0.75, radius=None):
        sp = s.shapes.add_shape(kind, IN(l), IN(t), IN(w), IN(h))
        sp.shadow.inherit = False
        if fill is None: sp.fill.background()
        else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line_c is None: sp.line.fill.background()
        else: sp.line.color.rgb = line_c; sp.line.width = PT(line_w)
        if radius is not None:
            try: sp.adjustments[0] = radius
            except Exception: pass
        sp.text_frame.paragraphs[0].text = ""
        return sp

    def rect(self, s, l, t, w, h, fill, line_c=None, line_w=0.75):
        return self._shape(s, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line_c, line_w)

    def card(self, s, l, t, w, h, fill=WHITE, line_c=LINE, line_w=1.25, radius=0.05):
        return self._shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill, line_c, line_w, radius)

    def oval(self, s, l, t, w, h, fill=LAV, line_c=PURPLE, line_w=1.25):
        return self._shape(s, MSO_SHAPE.OVAL, l, t, w, h, fill, line_c, line_w)

    def divider(self, s, l, t, w, fill=LINE, h=0.012):
        return self.rect(s, l, t, w, h, fill)

    def accent_bar(self, s, x, y, h, w=0.1, fill=PURPLE):
        return self.rect(s, x, y, w, h, fill)

    def logo(self, s, h=0.40, right=ML, top=0.34):
        w = h * LOGO_AR
        s.shapes.add_picture(self.logo_path, IN(DW - right - w), IN(top), height=IN(h))

    def pagenum(self, s, n):
        _, tf = self.box(s, DW - ML - 0.8, DH - 0.42, 0.8, 0.3)
        self.para(tf, {'text': str(n), 'size': 9, 'color': MUTE}, align=PP_ALIGN.RIGHT, first=True)

    # ----- the standard Braze leadline (purple header + gold subline + rule) -----
    def header(self, s, title, sub=None, kicker=None, tsize=23, logo=True, page=None):
        if logo: self.logo(s)
        if page is not None: self.pagenum(s, page)
        y = 0.5
        if kicker:
            _, tf = self.box(s, ML, y, CW - 2.4, 0.3)
            self.para(tf, {'text': kicker.upper(), 'size': 11, 'color': GOLD,
                           'bold': True, 'spc': 200}, first=True)
            y += 0.34
        _, tf = self.box(s, ML, y, CW - 0.2, 0.8)
        self.para(tf, {'text': title, 'size': tsize, 'color': PURPLE, 'bold': True},
                  first=True, line=1.0)
        y += 0.16 + tsize / 50.0
        if sub:
            _, tf = self.box(s, ML, y, CW, 0.5)
            self.para(tf, {'text': sub, 'size': 13, 'color': GOLD, 'bold': True},
                      first=True, line=1.05)
            y += 0.4
        self.rect(s, ML, y + 0.05, 1.3, 0.045, GOLD)
        return y + 0.3   # content top

    def footnote(self, s, text, y=7.02):
        _, tf = self.box(s, ML, y, CW, 0.42)
        self.para(tf, {'text': text, 'size': 7.5, 'color': MUTE, 'italic': True},
                  first=True, line=1.02)

    # ================= ARCHETYPE SLIDES =================
    def title(self, title, kicker="AI Decisioning", subtitle=None, confidential=True):
        s = self.blank()
        self.accent_bar(s, 0, 0, DH, w=0.16, fill=PURPLE)
        self.logo(s, h=0.46)
        if kicker:
            _, tf = self.box(s, ML + 0.2, 2.45, 9.5, 0.4)
            self.para(tf, {'text': kicker.upper(), 'size': 13, 'color': GOLD,
                           'bold': True, 'spc': 240}, first=True)
        _, tf = self.box(s, ML + 0.18, 2.9, 11.4, 1.6)
        self.para(tf, {'text': title, 'size': 40, 'color': PURPLE, 'bold': True},
                  first=True, line=1.0)
        self.rect(s, ML + 0.22, 4.35, 2.2, 0.06, GOLD)
        if subtitle:
            _, tf = self.box(s, ML + 0.2, 4.6, 10.0, 0.4)
            self.para(tf, {'text': subtitle, 'size': 15, 'color': SLATE}, first=True)
        if confidential:
            _, tf = self.box(s, ML + 0.2, 6.75, 6.0, 0.35)
            self.para(tf, {'text': 'Confidential', 'size': 10.5, 'color': MUTE,
                           'bold': True, 'spc': 60}, first=True)
        return s

    def section(self, title, subtitle=None, kicker=None):
        """Full leadline divider on lavender."""
        s = self.prs.slides.add_slide(self._blank)
        self.rect(s, 0, 0, DW, DH, LAV)
        self.accent_bar(s, 0, 0, DH, w=0.16, fill=PURPLE)
        self.logo(s, h=0.46)
        if kicker:
            _, tf = self.box(s, ML + 0.2, 2.35, 9.5, 0.4)
            self.para(tf, {'text': kicker.upper(), 'size': 13, 'color': GOLD,
                           'bold': True, 'spc': 220}, first=True)
        _, tf = self.box(s, ML + 0.18, 2.8, 11.4, 1.5)
        self.para(tf, {'text': title, 'size': 34, 'color': PURPLE, 'bold': True},
                  first=True, line=1.0)
        self.rect(s, ML + 0.22, 3.95, 2.2, 0.06, GOLD)
        if subtitle:
            _, tf = self.box(s, ML + 0.2, 4.2, 11.0, 0.6)
            self.para(tf, {'text': subtitle, 'size': 15, 'color': SLATE}, first=True, line=1.1)
        return s

    def two_column(self, title, left, right, sub=None, kicker=None, note=None, page=None):
        """left/right = (heading, [bullets]). Bullets get gold dots."""
        s = self.blank()
        cy = self.header(s, title, sub=sub, kicker=kicker, page=page)
        top = cy + 0.1; colw = (CW - 0.4) / 2; ch = 3.35 if not note else 3.15
        for i, (heading, bullets) in enumerate((left, right)):
            x = ML + i * (colw + 0.4)
            self.card(s, x, top, colw, ch, fill=WHITE)
            self.accent_bar(s, x, top, ch, fill=(PURPLE if i == 0 else GOLD))
            _, tf = self.box(s, x + 0.32, top + 0.25, colw - 0.6, 0.4)
            self.para(tf, {'text': heading, 'size': 15, 'color': PURPLE, 'bold': True}, first=True)
            _, tf = self.box(s, x + 0.32, top + 0.78, colw - 0.62, ch - 1.0)
            for j, b in enumerate(bullets):
                self.para(tf, [{'text': '•  ', 'size': 12, 'color': GOLD, 'bold': True},
                               {'text': b, 'size': 11.5, 'color': SLATE}],
                          first=(j == 0), line=1.16, after=8)
        if note:
            _, tf = self.box(s, ML, top + ch + 0.2, CW, 0.5)
            self.para(tf, note if isinstance(note, list) else
                      {'text': note, 'size': 11, 'color': SLATE}, first=True, line=1.15)
        return s

    def cards(self, title, items, sub=None, kicker=None, anchor=None, page=None):
        """items = [(num/label, heading, desc)]. Numbered cards in a row."""
        s = self.blank()
        cy = self.header(s, title, sub=sub, kicker=kicker, page=page)
        top = cy + 0.1; n = len(items); gap = 0.35
        cwm = (CW - (n - 1) * gap) / n; ch = 2.95 if anchor else 3.3
        for i, (num, heading, desc) in enumerate(items):
            x = ML + i * (cwm + gap)
            self.card(s, x, top, cwm, ch, fill=WHITE)
            self.rect(s, x, top, cwm, 0.12, GOLD)
            self.oval(s, x + 0.28, top + 0.35, 0.62, 0.62, fill=LAV, line_c=PURPLE)
            _, tf = self.box(s, x + 0.28, top + 0.36, 0.62, 0.6, anchor=MSO_ANCHOR.MIDDLE)
            self.para(tf, {'text': str(num), 'size': 22, 'color': PURPLE, 'bold': True},
                      align=PP_ALIGN.CENTER, first=True)
            _, tf = self.box(s, x + 0.3, top + 1.12, cwm - 0.6, 0.6)
            self.para(tf, {'text': heading, 'size': 15, 'color': PURPLE, 'bold': True}, first=True, line=1.0)
            _, tf = self.box(s, x + 0.3, top + 1.62, cwm - 0.58, 1.2)
            self.para(tf, {'text': desc, 'size': 11, 'color': SLATE}, first=True, line=1.2)
        if anchor:
            ay = top + ch + 0.22
            self.card(s, ML, ay, CW, 0.86, fill=LAV)
            self.accent_bar(s, ML, ay, 0.86, fill=PURPLE)
            _, tf = self.box(s, ML + 0.3, ay + 0.12, CW - 0.6, 0.66, anchor=MSO_ANCHOR.MIDDLE)
            lead, rest = anchor
            self.para(tf, [{'text': lead + '  ', 'size': 11.5, 'color': PURPLE, 'bold': True},
                           {'text': rest, 'size': 11.5, 'color': INK}], first=True, line=1.18)
        return s

    def stat_row(self, title, stats, sub=None, kicker=None, page=None, y=None, slide=None):
        """stats = [(value, label)]. A row of lavender stat chips. Reusable as a band."""
        s = slide or self.blank()
        if slide is None:
            cy = self.header(s, title, sub=sub, kicker=kicker, page=page)
            y = (cy + 0.2) if y is None else y
        n = len(stats); gap = 0.35; sw = (CW - (n - 1) * gap) / n
        for i, (val, lab) in enumerate(stats):
            x = ML + i * (sw + gap)
            self.card(s, x, y, sw, 0.92, fill=LAV, line_w=1.0)
            _, tf = self.box(s, x + 0.3, y + 0.14, sw - 0.6, 0.66, anchor=MSO_ANCHOR.MIDDLE)
            self.para(tf, [{'text': str(val) + '   ', 'size': 19, 'color': PURPLE, 'bold': True},
                           {'text': lab, 'size': 10.5, 'color': SLATE}], first=True, line=1.0)
        return s

    def table(self, title, headers, rows, sub=None, kicker=None, note=None,
              col_widths=None, cell_color=None, page=None, header_size=9,
              body_size=8.4, row_h=None, top=None):
        """Generic Braze table. rows = list of row-lists (strings).
        cell_color(ri, ci, value) -> RGBColor|None to tint specific text."""
        s = self.blank()
        cy = self.header(s, title, sub=sub, kicker=kicker, page=page)
        top = (cy + 0.05) if top is None else top
        nC = len(headers); nR = len(rows)
        total_h = (row_h * (nR + 1)) if row_h else min(4.4, 0.34 + 0.33 * nR)
        t = s.shapes.add_table(nR + 1, nC, IN(ML), IN(top), IN(CW), IN(total_h)).table
        t.first_row = False; t.horz_banding = False
        widths = col_widths or [CW / nC] * nC
        for i, w in enumerate(widths): t.columns[i].width = IN(w)
        for j, htext in enumerate(headers):
            c = t.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = PURPLE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = IN(0.07); c.margin_right = IN(0.05)
            c.margin_top = IN(0.02); c.margin_bottom = IN(0.02)
            self.para(c.text_frame, {'text': htext, 'size': header_size, 'color': WHITE, 'bold': True},
                      align=PP_ALIGN.LEFT if j else PP_ALIGN.CENTER, first=True)
        t.rows[0].height = IN(0.34)
        for ri, row in enumerate(rows, start=1):
            fill = WHITE if ri % 2 else LAV
            for ci, val in enumerate(row):
                c = t.cell(ri, ci); c.fill.solid(); c.fill.fore_color.rgb = fill
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.margin_left = IN(0.07); c.margin_right = IN(0.05)
                c.margin_top = IN(0.01); c.margin_bottom = IN(0.01)
                col = (cell_color(ri - 1, ci, val) if cell_color else None) or (INK if ci == 0 else SLATE)
                bold = ci == 0 or (cell_color is not None and cell_color(ri - 1, ci, val) is not None)
                self.para(c.text_frame, {'text': str(val), 'size': body_size, 'color': col, 'bold': bold},
                          align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER, first=True)
            t.rows[ri].height = IN((total_h - 0.34) / nR)
        if note: self.footnote(s, note, y=top + total_h + 0.08)
        return s

    def bar_chart(self, title, cats, values, top_labels=None, baseline=None,
                  baseline_label=None, sub=None, kicker=None, note=None,
                  page=None, subcats=None):
        """Vertical bar chart drawn as native shapes (full brand control).
        baseline: a value to draw a gold reference line at (e.g. break-even)."""
        s = self.blank()
        self.header(s, title, sub=sub, kicker=kicker, page=page)
        base_y = 5.0; px = ML + 0.05; pw = CW - 0.1; top_h = 2.45
        mx = max(values) or 1; n = len(values)
        pitch = pw / n; bw = pitch * 0.52
        if baseline is not None:
            by = base_y - (baseline / mx) * top_h
            self.rect(s, px, by, pw, 0.018, GOLD)
            if baseline_label:
                _, tf = self.box(s, px, by - 0.30, 4.2, 0.25)
                self.para(tf, {'text': baseline_label, 'size': 9, 'color': GOLD, 'bold': True}, first=True)
        for i, v in enumerate(values):
            cx = px + i * pitch + (pitch - bw) / 2
            h = (v / mx) * top_h
            if v <= 0:
                self.rect(s, cx, base_y - 0.03, bw, 0.03, LINE)
            else:
                col = PURPLE if (baseline is None or v >= baseline) else PERI
                self.rect(s, cx, base_y - h, bw, h, col)
                if top_labels:
                    _, tf = self.box(s, cx - 0.15, base_y - h - 0.26, bw + 0.3, 0.24)
                    self.para(tf, {'text': str(top_labels[i]), 'size': 8.5, 'color': INK, 'bold': True},
                              align=PP_ALIGN.CENTER, first=True)
            _, tf = self.box(s, cx - 0.2, base_y + 0.05, bw + 0.4, 0.42)
            self.para(tf, {'text': str(cats[i]), 'size': 9, 'color': INK, 'bold': True},
                      align=PP_ALIGN.CENTER, first=True)
            if subcats:
                self.para(tf, {'text': str(subcats[i]), 'size': 7.5, 'color': MUTE},
                          align=PP_ALIGN.CENTER, before=0)
        self.divider(s, px, base_y, pw)
        if note: self.footnote(s, note, y=6.95)
        return s

    def callout(self, s, l, t, w, h, lead, rest, fill=PURPLE, lead_c=GOLD, rest_c=WHITE, size=12.5):
        """A filled emphasis banner: bold gold lead-in + white body."""
        self.card(s, l, t, w, h, fill=fill, line_c=None, radius=0.05)
        _, tf = self.box(s, l + 0.35, t, w - 0.7, h, anchor=MSO_ANCHOR.MIDDLE)
        self.para(tf, [{'text': lead + '    ', 'size': size, 'color': lead_c, 'bold': True, 'spc': 100},
                       {'text': rest, 'size': size, 'color': rest_c, 'bold': True}], first=True, line=1.05)

    def closing(self, headline, subline=None, footer="BrazeAI Decisioning Studio™  ·  Confidential"):
        s = self.prs.slides.add_slide(self._blank)
        self.rect(s, 0, 0, DW, DH, LAV)
        self.accent_bar(s, 0, 0, DH, w=0.16, fill=PURPLE)
        self.rect(s, 0, DH - 0.16, DW, 0.16, GOLD)
        self.logo(s, h=0.5)
        _, tf = self.box(s, ML + 0.2, 2.85, 11.5, 1.4)
        self.para(tf, {'text': headline, 'size': 34, 'color': PURPLE, 'bold': True}, first=True, line=1.0)
        self.rect(s, ML + 0.22, 3.95, 2.2, 0.06, GOLD)
        if subline:
            _, tf = self.box(s, ML + 0.2, 4.2, 11.0, 0.5)
            self.para(tf, {'text': subline, 'size': 15, 'color': SLATE}, first=True)
        if footer:
            _, tf = self.box(s, ML + 0.2, 6.75, 8.0, 0.35)
            self.para(tf, {'text': footer, 'size': 10.5, 'color': MUTE}, first=True)
        return s

    def save(self, path):
        self.prs.save(path)
        return path
